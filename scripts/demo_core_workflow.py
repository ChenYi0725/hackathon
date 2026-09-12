"""Build an isolated local demo from synthetic inputs, never invoke AWS."""
import argparse
import io
import json
from pathlib import Path
from openpyxl import Workbook
from app.bootstrap import build_service
from app.domain.models import Case
from app.infrastructure.settings import Settings


def rules():
    return dict(name='合成流程驗收基準（非真實估價規則）', version='synthetic-1', locality='合成區', land_use='合成用地',
        source='scripts/demo_core_workflow.py；人工固定合成矩陣，僅用於流程測試',
        direction='列：比準地；欄：比較標的；矩陣值為百分點',
        rules=[dict(id='width', name='寬度', group='道路', unit='m', scope='individual', source_page=1,
                    bands=[dict(label='窄',low=0,high=10),dict(label='寬',low=10,high=None)],matrix=[[0,-2],[2,0]]),
               dict(id='region', name='區域條件', group='區域', unit='', scope='regional', source_page=1,
                    bands=[dict(label='甲',values=['甲']),dict(label='乙',values=['乙'])],matrix=[[0,-1],[1,0]])])


def synthetic_workbook(mode='正常'):
    book=Workbook();survey=book.active;survey.title='勘查'
    for cell,value in {'A1':'表3 地價區段勘查表','A3':'年期','B3':'2026-09-01','G3':'S1','D11':'主要道路','J11':12,'F12':8,'H6':0,'H7':'無','H8':'不適用','J12':'=J11+1'}.items():survey[cell]=value
    regional=book.create_sheet('區域因素')
    for cell,value in {'A1':'影響地價區域因素分析明細表','A3':'地價區段號','C3':'S1','E3':'S2','G4':'修正百分比','C5':'甲','E5':'甲','G5':0,'E42':0}.items():regional[cell]=value
    comparison=book.create_sheet('比較法')
    for cell,value in {'A1':'比較法調查估價表','C8':'區域因素調整百分率','G8':'S2','J8':0,'D10':12,'G10':8,'J10':2,'G5':100,'G7':100}.items():comparison[cell]=value
    for cell,value in {'J6':0,'J28':2,'J30':2,'J31':102,'J32':100,'A28':'原填個別因素合計','A30':'原填絕對值合計','A31':'原填試算價格','A32':'原填權重'}.items():comparison[cell]=value
    if mode=='錯誤':comparison['J10']=9
    if mode=='缺資料':survey['J11']=None
    for sheet in book:
        sheet.freeze_panes='B4'
        for column in 'ABCDEFGHIJ':sheet.column_dimensions[column].width=18
    stream=io.BytesIO();book.save(stream);return stream.getvalue()


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--data-dir',type=Path,default=Path('.analysis/core-demo'))
    parser.add_argument('--fixture-only',type=Path)
    args=parser.parse_args()
    if args.fixture_only:
        args.fixture_only.parent.mkdir(parents=True,exist_ok=True)
        args.fixture_only.write_bytes(synthetic_workbook());return
    if (args.data_dir/'review.sqlite3').exists():
        raise SystemExit('Demo 資料庫已存在；請指定另一個 --data-dir，不覆寫既有案件。')
    service=build_service(Settings(data_dir=args.data_dir,ai_enabled=False))
    draft=service.create_ruleset(rules())
    published=service.publish_ruleset(draft['id'],'核對合成驗收矩陣，非真實規則',True,True,'2026-01-01','2026-12-31')
    outputs=[]
    for mode in ('正常','錯誤','缺資料','外部失敗'):
        case=Case.model_validate(dict(title='合成 Demo · '+mode,case_number='SYNTHETIC-'+mode,
            valuation_date='2026-09-01',locality='合成區',land_use='合成用地',ruleset_id=published['id'],
            subject_name='合成比準地',subject_section='S1',comparable_name='合成比較標的',comparable_section='S2',
            factors=[dict(id='width',subject=None if mode=='缺資料' else '12',comparable='8',entered_rate=9 if mode=='錯誤' else 2),
                     dict(id='region',subject='甲',comparable='甲',subject_grade='甲',comparable_grade='甲',entered_rate=0)],
            totals=dict(normal_price=100,time_rate=0,adjusted_price=100,regional_detail=0,regional_carried=0,individual=2,absolute=2,trial_price=102,weight=100)))
        created=service.save_case(case,new=True)
        attached=service.workflow.upload(created['case']['id'],created['case']['revision'],synthetic_workbook(mode),'合成三種書表.xlsx')
        docid=attached['document_id']
        mappings={'factors.width.subject':('勘查','J11'),'factors.width.comparable':('比較法','G10'),'factors.width.entered_rate':('比較法','J10'),
                  'factors.region.subject':('區域因素','C5'),'factors.region.comparable':('區域因素','E5'),'factors.region.entered_rate':('區域因素','G5'),
                  'totals.regional_detail':('區域因素','E42'),'totals.regional_carried':('比較法','J8'),'totals.individual':('比較法','J28'),
                  'totals.absolute':('比較法','J30'),'totals.time_rate':('比較法','J6'),'totals.normal_price':('比較法','G5'),
                  'totals.adjusted_price':('比較法','G7'),'totals.trial_price':('比較法','J31'),'totals.weight':('比較法','J32')}
        for target,(sheet,cell) in mappings.items():
            attached=service.workflow.apply_cell(created['case']['id'],attached['case']['revision'],docid,sheet,cell,target)
        case=Case.model_validate(attached['case'])
        for factor in case.factors:factor.confirmed=True
        case.totals_confirmed=True
        saved=service.save_case(case)
        if mode=='外部失敗':
            service.workflow.lookup(case.id,saved['case']['revision'],'width',dict(mode='mock',scenario='timeout'))
            saved=service.get_case(case.id)
        data,_,_=service.workflow.export(case.id,saved['case']['revision'],saved['run']['id'],'bundle')
        (args.data_dir/f'demo-{mode}.zip').write_bytes(data)
        outputs.append(dict(case_id=case.id,mode=mode,complete=saved['review']['complete'],counts=saved['review']['counts']))
    (args.data_dir/'synthetic-forms.xlsx').write_bytes(synthetic_workbook())
    print(json.dumps(outputs,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
