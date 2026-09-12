"""Bounded Excel inspection and existing Paddle PDF port; no formula execution."""
import hashlib
import io
import zipfile
from datetime import date, datetime
from decimal import Decimal
from openpyxl import load_workbook


class CaseDocumentReader:
    def __init__(self, pdf):
        self.pdf = pdf

    def read(self, data, name):
        if len(data) > 20 * 1024 * 1024:
            raise ValueError('文件上限 20 MB。')
        if data.startswith(b'%PDF-'):
            pages = self.pdf.read(data)
            if not pages or not any(p.get('text', '').strip() for p in pages):
                raise ValueError('PDF 未辨識出文字，請檢查掃描品質。')
            return dict(kind='pdf', pages=pages, warnings=['PDF 已由 PaddleOCR 辨識；目前自動欄位解析限既有金山版型，其餘需人工核對。'])
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ValueError('僅接受有效 PDF 或 XLSX 文件。')
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = z.namelist()
            if len(names) > 2000 or sum(i.file_size for i in z.infolist()) > 30 * 1024 * 1024:
                raise ValueError('Excel 解壓內容超過限制。')
            if any('vbaproject' in n.lower() or 'externallinks/' in n.lower() for n in names):
                raise ValueError('不接受巨集或外部連結。')
        try:
            book = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
            values = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        except Exception:
            raise ValueError('無法解析 Excel。') from None
        pages, warnings = [], []
        sha = hashlib.sha256(data).hexdigest()
        try:
            if len(book.worksheets) > 20:
                raise ValueError('最多接受 20 個工作表。')
            for index, sheet in enumerate(book.worksheets, 1):
                if sheet.max_row is None or sheet.max_column is None:
                    raise ValueError('Excel 缺少工作表範圍資訊，請以試算表軟體重新儲存為 XLSX。')
                if sheet.max_row > 2000 or sheet.max_column > 100:
                    raise ValueError('工作表範圍超過 2000 列／100 欄。')
                if sheet.sheet_state != 'visible':
                    warnings.append('含隱藏工作表：' + sheet.title + '；僅本機保存，不自動送雲端。')
                title = str(sheet['A1'].value or '')
                labels = ' '.join(str(sheet.cell(r, c).value or '') for r in range(1, 13) for c in range(1, 11))
                kind = ('survey' if '地價區段勘查表' in title and '主要道路' in labels and '年期' in labels else
                        'regional' if '影響地價區域因素' in title and '修正百分比' in labels and '地價區段號' in labels else
                        'comparison' if '比較法調查估價表' in title and '區域因素調整百分率' in labels else 'unsupported')
                if kind == 'unsupported':
                    warnings.append(sheet.title + '：未知版型；保存原值供人工映射，不宣稱解析完成。')
                cells = {}
                cached = values[sheet.title]
                for row in sheet:
                    for cell in row:
                        if cell.value is None:
                            continue
                        formula = str(cell.value) if cell.data_type == 'f' else None
                        value = cached[cell.coordinate].value if formula else cell.value
                        value_type = 'date' if isinstance(value, (date, datetime)) else 'scalar'
                        if isinstance(value, (date, datetime)):
                            value = value.isoformat()
                        if isinstance(value, (int, float)) and not isinstance(value, bool) and '%' in cell.number_format:
                            value = str(Decimal(str(value)) * 100)
                        cells[cell.coordinate] = dict(value=value, value_type=value_type, formula=formula, number_format=cell.number_format,
                            presence='blank' if value in (None, '') else 'none' if value == '無' else
                            'not_applicable' if value == '不適用' else 'value')
                text = '\n'.join(f'{k}: {v["value"] if v["value"] is not None else "[空白]"}' +
                                 (f' [公式 {v["formula"]}]' if v['formula'] else '') for k, v in cells.items())
                pages.append(dict(page=index, text=text, sheet=sheet.title, kind=kind, method='xlsx',
                                  sha256=sha, hidden=sheet.sheet_state != 'visible', cells=cells,
                                  max_row=sheet.max_row, max_column=sheet.max_column))
        finally:
            book.close()
            values.close()
        return dict(kind='xlsx', pages=pages, warnings=warnings + ['Excel 公式不執行；快取不代表已驗證。先指定儲存格及目標欄位，再人工確認。'])
