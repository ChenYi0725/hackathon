"""Fill the supplied worksheets and render the same values to PDF.

Only presentation mappings live here. No valuation arithmetic or Excel formulas
are executed. Missing fields stay explicit; the original templates are read-only.
"""
import io
import hashlib
import os
import warnings
from copy import copy
from functools import lru_cache
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import BadZipFile

from app.application.export_contracts import ExportArtifact, ExportUnavailable

STATUS = {'pass': '通過', 'error': '疑似錯誤', 'pending': '待確認', 'missing': '資料不足'}
TEMPLATES = {'3': ('表3*.xlsx', '表3區段勘查表', 46, 22),
             '4': ('表4*.xlsx', '表4比較法調查估價表', 37, 18),
             '5': ('表5*.xlsx', '表5-1區域因素明細表(住)', 44, 13)}
INDIVIDUAL = dict(zip(range(9, 28), (
    'area', 'width', 'depth', 'shape', 'frontage', 'terrain', 'road_type',
    'road_width', 'school', 'market', 'park', 'station', 'business', 'nuisance',
    'parking', 'zoning', 'coverage', 'far', 'restriction')))
REGIONAL = {5: 'r_plan', 6: 'r_zoning', 7: 'r_coverage', 8: 'r_far',
    9: 'r_ban', 10: 'r_restriction', 12: 'r_road_width', 13: 'r_avg_width',
    14: 'r_station', 15: 'r_bus', 16: 'r_interchange', 17: 'r_development',
    19: 'r_sunlight', 20: 'r_landscape', 21: 'r_slope', 22: 'r_drainage',
    23: 'r_terrain', 25: 'r_improvement', 27: 'r_school', 28: 'r_market',
    29: 'r_park', 30: 'r_tourism', 31: 'r_parking', 32: 'r_service',
    34: 'r_utility', 35: 'r_funeral', 36: 'r_waste', 38: 'r_pollution', 40: 'r_other'}
SURVEY = {'H4': 'r_plan', 'H5': 'r_zoning', 'H6': 'r_coverage', 'H7': 'r_far',
    'H8': 'r_ban', 'H9': 'r_restriction', 'J11': 'r_road_width', 'G12': 'r_avg_width',
    'J19': 'r_interchange', 'I23': 'r_development', 'I24': 'r_sunlight',
    'I25': 'r_landscape', 'I26': 'r_slope', 'I27': 'r_drainage', 'I28': 'r_terrain',
    'I29': 'r_wind', 'I30': 'r_soil', 'R38': 'r_customers', 'R39': 'r_shops'}


def text(value):
    return '待補' if value is None or value == '' else str(value)


def put(ws, address, value):
    """Write literal text, including strings starting '='; never spreadsheet code."""
    from openpyxl.styles import Alignment
    from openpyxl.cell.cell import MergedCell
    cell = ws[address]
    if isinstance(cell, MergedCell):
        raise ExportUnavailable('書表模板合併儲存格與填值契約不符。')
    numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
    cell.value = value if numeric else text(value)
    cell.data_type = 'n' if numeric else 's'
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    font = copy(cell.font)
    font.color = '8A5200' if '待' in str(cell.value) or '未提供' in str(cell.value) else '174E44'
    cell.font = font


@lru_cache(maxsize=1)
def pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.ttfonts import TTFError
    candidates = [os.getenv('PDF_FONT_PATH'),
                  str(Path(os.getenv('WINDIR', 'C:/Windows')) / 'Fonts/msjh.ttc'),
                  '/usr/share/fonts/truetype/arphic/uming.ttc']
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            try:
                pdfmetrics.registerFont(TTFont('ExportChinese', candidate))
            except TTFError as error:
                raise ExportUnavailable('PDF_FONT_PATH 字型無法嵌入；請使用 TrueType 輪廓的繁體中文字型。') from error
            return 'ExportChinese'
    raise ExportUnavailable('PDF 中文字型未設定；請以 PDF_FONT_PATH 指定支援繁體中文的 TrueType 字型。')


class TemplateFormRenderer:
    def __init__(self, template_dir):
        self.template_dir = Path(template_dir)

    def workbook(self, case, result, rules, number, generated_at):
        from openpyxl import load_workbook
        pattern, sheet_name, rows, cols = TEMPLATES[number]
        paths = sorted(self.template_dir.glob(pattern))
        if len(paths) != 1:
            raise ExportUnavailable(f'請在 FORM_TEMPLATE_DIR 放入唯一的表{number} Excel 模板。')
        source = paths[0].read_bytes()
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            try:
                wb = load_workbook(io.BytesIO(source), keep_links=False)
            except BadZipFile as error:
                raise ExportUnavailable(f'表{number} 模板不是有效的 XLSX 檔案。') from error
        if sheet_name not in wb:
            raise ExportUnavailable(f'表{number} 模板缺少指定工作表。')
        for sheet in list(wb):
            if sheet.title != sheet_name:
                wb.remove(sheet)
        ws = wb[sheet_name]
        ws.sheet_state = 'visible'
        from openpyxl.utils import get_column_letter
        if not any('A1' in merged for merged in ws.merged_cells.ranges):
            ws.merge_cells(f'A1:{get_column_letter(cols)}1')
        # Remove template example formulas, names and links into hidden examples.
        wb.defined_names.clear()
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    put(ws, cell.coordinate, '待確認')
                if cell.hyperlink:
                    cell.hyperlink = None
        factors = {f.id: f for f in case.factors}
        def factor(fid, side):
            f = factors.get(fid)
            return None if f is None else getattr(f, side)
        note = (f'{case.title}｜案件 {case.id}｜版本 {case.revision}｜{case.locality} {case.land_use}｜'
                f'基準 {rules["id"]}/{rules["version"]}｜原填值草稿；待補／未確認值不得視為已核准')
        if number == '3':
            other = wb.copy_worksheet(ws)
            ws.title, other.title = '表3-比準地', '表3-比較標的1'
            for sheet, side in [(ws, 'subject'), (other, 'comparable')]:
                put(sheet, 'B3', case.valuation_date)
                put(sheet, 'G3', getattr(case, side + '_section'))
                put(sheet, 'L3', f'{case.locality} {getattr(case, side + "_name")}；範圍待核對')
                for address, fid in SURVEY.items():
                    put(sheet, address, factor(fid, side))
                # Facility type/name is not represented by the legacy case. Do not
                # turn a generic station distance into a high-speed-rail distance.
                for address in ['F11', 'F13', 'F14', 'F15', 'F16', 'F17', 'G18',
                    'F19', 'I20', 'I21', 'I22', 'E31', 'E32', 'E33', 'E34',
                    'F35', 'F36', 'F37', 'F38', 'F39', 'F40', 'F41', 'F42', 'F43', 'F44',
                    'R4', 'R6', 'R8', 'R10', 'R11', 'R12', 'R14', 'R16',
                    'R18', 'R19', 'R20', 'R21', 'R22', 'R23', 'R24', 'R25',
                    'R26', 'R27', 'R28', 'R29', 'S30', 'S32', 'S34', 'S36',
                    'R42', 'R43', 'Q44']:
                    put(sheet, address, '待補（見填值明細）')
                put(sheet, 'L40', f'版本 {case.revision}｜草稿；完整資料見明細')
        elif number == '4':
            put(ws, 'L1', case.valuation_date)
            put(ws, 'O1', '案號：' + text(case.case_number))
            for address, value in {'F2': case.subject_name, 'J2': case.comparable_name,
                'D4': case.subject_name, 'G4': case.comparable_name,
                'D8': case.subject_section, 'G8': case.comparable_section,
                'G5': case.totals.normal_price, 'J6': case.totals.time_rate,
                'G7': case.totals.adjusted_price, 'J8': case.totals.regional_carried,
                'G29': case.totals.individual, 'G30': case.totals.absolute,
                'G31': case.totals.trial_price, 'I31': case.totals.weight,
                'G32': '待確認（未有獨立比準地比較價格欄位）', 'G6': '交易日期待補',
                'D33': case.notes, 'D34': note}.items():
                put(ws, address, value)
            for row, fid in INDIVIDUAL.items():
                for col, side in [('D', 'subject'), ('G', 'comparable'), ('J', 'entered_rate')]:
                    put(ws, f'{col}{row}', factor(fid, side))
            for col in ('K', 'O'):
                for row in range(4, 30):
                    put(ws, f'{col}{row}', '未提供標的')
            for address in ('N2', 'R2', 'K30', 'O30', 'K31', 'O31', 'D28', 'G28', 'J28'):
                put(ws, address, '未提供')
        else:
            put(ws, 'A2', '案號：' + text(case.case_number))
            put(ws, 'G2', case.comparable_name)
            put(ws, 'C3', case.subject_section)
            put(ws, 'E3', case.comparable_section)
            residential = '住宅' in case.land_use
            for row, fid in REGIONAL.items():
                for col, side in [('C', 'subject_grade'), ('E', 'comparable_grade'), ('G', 'entered_rate')]:
                    put(ws, f'{col}{row}', factor(fid, side) if residential else '基準不適用')
                for col in ('H', 'J', 'K', 'M'):
                    put(ws, f'{col}{row}', '未提供')
            for row in (11, 18, 24, 26, 33, 37, 39, 41):
                put(ws, f'E{row}', '小計待確認')
                put(ws, f'H{row}', '未提供')
                put(ws, f'K{row}', '未提供')
            put(ws, 'E42', case.totals.regional_detail if residential else '基準不適用')
            put(ws, 'H42', '未提供')
            put(ws, 'K42', '未提供')
            put(ws, 'B42', '各主要項目總修正數')
            put(ws, 'C43', case.subject_name)
            put(ws, 'E43', case.comparable_name)
            put(ws, 'C44', note + ('；此住宅表不適用商業案件' if not residential else ''))
        for sheet in wb:
            sheet.print_area = f'A1:{get_column_letter(cols)}{rows}'
            sheet.page_setup.orientation = 'portrait' if number in ('3', '5') else 'landscape'
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 1
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
        detail = wb.create_sheet('填值與審核明細')
        detail.append(['模板', paths[0].name, 'SHA256', hashlib.sha256(source).hexdigest()])
        for row in self.detail_rows(case, result, rules, generated_at):
            detail.append([text(v) for v in row])
        for row in detail:
            for cell in row:
                cell.data_type = 's'
        return wb

    @staticmethod
    def detail_rows(case, result, rules, generated_at):
        yield ['案件', case.title, '版本', case.revision]
        yield ['基準', rules['id'], '基準版本', rules['version']]
        yield ['產出時間', generated_at, '狀態', '原填值草稿；詳見審核結果']
        yield ['案件備註', case.notes]
        yield ['限制', '目前一筆比較標的；其他標的未提供。表3設施名稱等缺值待補。']
        yield ['因素', '比準地原填', '比較標的原填', '原填修正率', '確認', '來源頁', '引用／備註']
        names = {r['id']: r['name'] for r in rules['rules']}
        for f in case.factors:
            yield [names.get(f.id, f.id), f.subject, f.comparable, f.entered_rate,
                   '已確認' if f.confirmed else '待確認', f.evidence.page, f.evidence.quote + ' ' + f.note]
        yield ['檢核項目', '狀態', '原填值', '預期值', '說明', '原文頁', '基準頁']
        for r in result['checks']:
            yield [r['title'], STATUS[r['status']], r.get('actual'), r.get('expected'),
                   r['message'], r.get('page'), r.get('rule_page')]

    def render(self, case, result, rules, kind, generated_at):
        if kind == 'report-pdf':
            data = self.report_pdf(case, result, rules, generated_at)
            extension = 'pdf'
        else:
            parts = kind.split('-')
            if len(parts) != 2 or parts[0] not in ('table3', 'table4', 'table5') or parts[1] not in ('xlsx', 'pdf'):
                raise KeyError(kind)
            number, extension = parts[0][-1], parts[1]
            wb = self.workbook(case, result, rules, number, generated_at)
            if extension == 'xlsx':
                stream = io.BytesIO()
                wb.save(stream)
                data = stream.getvalue()
            else:
                data = self.workbook_pdf(wb, case, generated_at, TEMPLATES[number][2:])
        media = 'application/pdf' if extension == 'pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        return ExportArtifact(data, f'{kind.rsplit("-", 1)[0]}-{case.id}-r{case.revision}.{extension}', media, case.revision)

    @staticmethod
    def report_pdf(case, result, rules, generated_at):
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle, Spacer
        font = pdf_font()
        style = ParagraphStyle('Chinese', fontName=font, fontSize=9, leading=13, wordWrap='CJK')
        p = lambda value: Paragraph(escape(text(value)).replace('\n', '<br/>'), style)
        stream = io.BytesIO()
        doc = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=28, rightMargin=28,
                                topMargin=28, bottomMargin=28, title=case.title)
        rows = [[p(v) for v in ['檢核項目', '狀態', '原填值', '預期值', '依據與說明']]]
        for r in result['checks']:
            rows.append([p(v) for v in [r['title'], STATUS[r['status']], r.get('actual'), r.get('expected'),
                f'{r["message"]}\n原文 p.{text(r.get("page"))}／基準 p.{text(r.get("rule_page"))}']])
        table = LongTable(rows, colWidths=[132, 58, 65, 65, 465], repeatRows=1, splitInRow=1)
        table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .3, '#9DAFA8'),
                                  ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                  ('BACKGROUND', (0, 0), (-1, 0), '#EAF3EF')]))
        doc.build([p('地衡｜逐項審查結果'), p(case.title),
            p(f'案件 {case.id}｜版本 {case.revision}｜基準 {rules["id"]}/{rules["version"]}'),
            p(f'比準地 {case.subject_name}／比較標的 {case.comparable_name}｜基準日 {case.valuation_date}'),
            p('輔助審查草稿；' + ('全部檢核通過' if result['complete'] else '尚有疑點或待確認項目')),
            p(case.notes), p(generated_at), Spacer(1, 12), table])
        return stream.getvalue()

    @staticmethod
    def workbook_pdf(wb, case, generated_at, size):
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.lib.pagesizes import A3, landscape
        from reportlab.pdfbase.pdfmetrics import stringWidth
        from openpyxl.cell.cell import MergedCell
        from openpyxl.utils import get_column_letter
        font = pdf_font()
        stream = io.BytesIO()
        page_w, page_h = landscape(A3)
        canvas = Canvas(stream, pagesize=(page_w, page_h))
        canvas.setTitle(case.title)
        rows, cols = size
        overflow = []
        for ws in list(wb)[:-1]:
            page_w, page_h = A3 if ws.page_setup.orientation == 'portrait' else landscape(A3)
            canvas.setPageSize((page_w, page_h))
            widths = [max(12, ws.column_dimensions[get_column_letter(c)].width * 5.2) for c in range(1, cols + 1)]
            heights = [max(15, ws.row_dimensions[r].height or 18) for r in range(1, rows + 1)]
            scale = min((page_w - 48) / sum(widths), (page_h - 80) / sum(heights))
            xs, ys = [24], [page_h - 48]
            for width in widths: xs.append(xs[-1] + width * scale)
            for height in heights: ys.append(ys[-1] - height * scale)
            merges = {(m.min_row, m.min_col): (m.max_row, m.max_col) for m in ws.merged_cells.ranges}
            canvas.setFont(font, 9)
            canvas.drawString(24, page_h - 25, f'{ws.title}｜案件版本 {case.revision}｜原填值草稿；待補欄位詳見附錄')
            for row in ws.iter_rows(max_row=rows, max_col=cols):
                for cell in row:
                    if isinstance(cell, MergedCell): continue
                    r, c = cell.row, cell.column
                    end_r, end_c = merges.get((r, c), (r, c))
                    end_r, end_c = min(end_r, rows), min(end_c, cols)
                    x, y, width, height = xs[c-1], ys[end_r], xs[end_c]-xs[c-1], ys[r-1]-ys[end_r]
                    canvas.setStrokeColorRGB(.65, .69, .67)
                    canvas.setLineWidth(.25)
                    if any(side and side.style for side in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom)):
                        canvas.rect(x, y, width, height)
                    if cell.value is None: continue
                    value = str(cell.value)
                    font_size = min(9, max(5, (cell.font.sz or 10) * scale))
                    while True:
                        lines, line = [], ''
                        for char in value:
                            if char == '\n' or stringWidth(line + char, font, font_size) > width - 4:
                                lines.append(line); line = '' if char == '\n' else char
                            else: line += char
                        lines.append(line)
                        if len(lines) * font_size * 1.1 <= height - 2 or font_size <= 3: break
                        font_size -= .25
                    if len(lines) * font_size * 1.1 > height - 2:
                        overflow.append([ws.title, cell.coordinate, value])
                        lines = ['全文見附錄']
                        font_size = min(5, max(2, (width - 4) / 5))
                    canvas.setFillColorRGB(.10, .19, .16)
                    canvas.setFont(font, font_size)
                    for i, line in enumerate(lines):
                        canvas.drawString(x+2, y+height-font_size-1-i*font_size*1.1, line)
            canvas.setFont(font, 8)
            canvas.drawString(24, 22, f'基準 {case.ruleset_id}｜{generated_at}｜缺值未補零；簽章欄由承辦人填寫')
            canvas.showPage()
        # Include all input values and review statuses, even those that do not
        # have a corresponding cell in the selected template.
        from pypdf import PdfReader, PdfWriter
        canvas.save()
        writer = PdfWriter()
        writer.append(PdfReader(io.BytesIO(stream.getvalue())))
        # Render the literal detail rows as a paginated appendix.
        from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle
        from reportlab.lib.styles import ParagraphStyle
        appendix = io.BytesIO()
        style = ParagraphStyle('Detail', fontName=font, fontSize=8, leading=11, wordWrap='CJK')
        data = [[Paragraph(escape(text(c.value)), style) for c in row] for row in wb.worksheets[-1]]
        for row in overflow:
            data.append([Paragraph(escape(text(v)), style) for v in row] + [''] * 4)
        table = LongTable(data, colWidths=[135, 150, 150, 85, 180, 60, 340], splitInRow=1)
        table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .3, '#AAAAAA'), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
        SimpleDocTemplate(appendix, pagesize=landscape(A3), leftMargin=35, rightMargin=35).build([table])
        writer.append(PdfReader(io.BytesIO(appendix.getvalue())))
        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()
