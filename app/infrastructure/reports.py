"""Render only an immutable review snapshot; no valuation calculations here."""
import io
import json
import zipfile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from app.application.ports import ExtractionUnavailable

LABELS = {'pass': '通過', 'error': '疑似錯誤', 'pending': '待確認', 'missing': '資料不足'}


class SnapshotRenderer:
    def pdf(self, run, decisions):
        # Reuse the existing report layout and font handling added by the form-export work.
        from app.domain.models import Case
        from app.infrastructure.form_exports import TemplateFormRenderer
        from app.application.export_contracts import ExportUnavailable
        result = dict(run['review'], run_id=run['id'], dispositions=decisions)
        try:
            return TemplateFormRenderer.report_pdf(Case.model_validate(run['case']), result, run['rules'], run['generated_at'])
        except ExportUnavailable as error:
            raise ExtractionUnavailable(str(error)) from error

    def xlsx(self, run, decisions):
        book = Workbook()
        sheet = book.active
        sheet.title = '已確認勘查資料'
        sheet.append(['案件 ID', run['case_id'], '案件版本', run['case_revision']])
        sheet.append(['基準 ID', run['rules']['id'], '基準版本', run['rules']['version']])
        sheet.append(['因素', '比準地', '比較標的', '原填修正率', '確認狀態', '來源'])
        names = {r['id']: r['name'] for r in run['rules']['rules']}
        def safe(v):
            return "'" + v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@')) else v
        all_factors = [('primary', f) for f in run['case']['factors']]
        for comparison in run['case'].get('additional_comparisons', []):
            all_factors.extend((comparison['id'], f) for f in comparison['factors'])
        sheet['G3'] = '比較標的 ID'
        for comparison_id, f in all_factors:
            sheet.append([safe(names.get(f['id'], f['id'])), safe(f['subject']) if f['confirmed'] else '待確認',
                          safe(f['comparable']) if f['confirmed'] else '待確認',
                          f['entered_rate'] if f['confirmed'] else '待確認',
                          '已確認' if f['confirmed'] else '未確認（原值見快照）', safe(' / '.join(str(f['evidence'].get(k) or '') for k in ('document_id','sheet','cell','page','method'))), comparison_id])
        totals = book.create_sheet('計算欄位')
        totals.append(['欄位', '原填值', '確認狀態'])
        for k, v in run['case']['totals'].items():
            totals.append([k, v, '已確認' if run['case']['totals_confirmed'] else '待確認'])
        for comparison in run['case'].get('additional_comparisons', []):
            for k, v in comparison['totals'].items():
                totals.append([comparison['id'] + ':' + k, v, '已確認' if comparison['totals_confirmed'] else '待確認'])
        review = book.create_sheet('審查摘要')
        review.append(['項目', '狀態', '原值', '預期值', '規則版本', '計算及來源'])
        for r in run['review']['checks']:
            review.append([safe(r['title']), LABELS[r['status']], r.get('actual'), r.get('expected'),
                           run['rules']['version'], safe(r['message'])])
        human = book.create_sheet('人工處置')
        human.append(['項目', '接受或拒絕', '原因', '時間', '適用結果'])
        for d in decisions:
            human.append([safe(d['check_id']), d['decision'], safe(d['reason']), d['at'],
                          '目前' if d['run_id'] == run['id'] else '歷史（不適用本結果）'])
        for ws in book:
            ws.freeze_panes = 'A2'
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill('solid', fgColor='E8F2EE')
            for col in 'ABCDEFG':
                ws.column_dimensions[col].width = 24 if col != 'F' else 65
            for row in ws:
                for cell in row:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                ws.row_dimensions[row[0].row].height = 45
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
        output = io.BytesIO()
        book.save(output)
        return output.getvalue()

    def render(self, run, decisions, kind):
        if kind == 'pdf':
            return self.pdf(run, decisions), 'application/pdf', 'review.pdf'
        if kind == 'xlsx':
            return self.xlsx(run, decisions), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'review.xlsx'
        if kind != 'bundle':
            raise ValueError('不支援的匯出類型。')
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr('review.pdf', self.pdf(run, decisions))
            z.writestr('confirmed-forms.xlsx', self.xlsx(run, decisions))
            z.writestr('snapshot.json', json.dumps(dict(run=run, dispositions=decisions), ensure_ascii=False, indent=2))
        return output.getvalue(), 'application/zip', 'review-bundle.zip'
