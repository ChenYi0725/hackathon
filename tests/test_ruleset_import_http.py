import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.infrastructure.settings import Settings
from app.interfaces.http import create_app
from tests.test_ruleset_imports import extraction_result


class WorkflowPdf:
    def read(self, data):
        if data == b'question-pdf':
            text = (
                '估價基準日：2026-09-12\n'
                '比準地：測試市甲區幸福路1號\n'
                '比較標的：測試市甲區和平路2號\n'
                '測試數值 85% 65% 1.25%\n'
            )
        else:
            text = '測試市甲區測試用地影響地價區域因素評價基準明細表'
        return [{
            'page': 1,
            'text': text,
            'lines': [],
            'width': 1000,
            'height': 1400,
            'method': 'test',
        }]


class WorkflowExtractor:
    def extract(self, pages, *, source_name, expected_locality):
        return extraction_result(expected_locality)


def test_ruleset_upload_confirmation_question_fill_and_excel(tmp_path, monkeypatch):
    monkeypatch.setenv('SEED_EXAMPLES', 'false')
    app = create_app(
        Settings(data_dir=tmp_path, ai_enabled=False),
        pdf=WorkflowPdf(),
        ruleset_extractor=WorkflowExtractor(),
    )
    with TestClient(app) as client:
        draft_response = client.post(
            '/api/ruleset-imports?expected_locality=測試市甲區&name=測試基準.pdf',
            content=b'ruleset-pdf',
            headers={'Content-Type': 'application/pdf'},
        )
        assert draft_response.status_code == 200, draft_response.text
        draft = draft_response.json()
        assert draft['requires_confirmation'] is True
        assert draft['candidates'][0]['locality'] == '測試市甲區'

        confirmation = client.post('/api/ruleset-imports/confirm', json={
            'document_id': draft['document_id'],
            'candidate': draft['candidates'][0],
            'valid_from': '2026-01-01',
            'valid_to': '2026-12-31',
            'matrix_direction': 'benchmark_row_target_column',
            'confirmed': True,
        })
        assert confirmation.status_code == 200, confirmation.text
        ruleset = confirmation.json()['ruleset']
        sources = client.get(
            f'/api/rulesets/{ruleset["id"]}/evidence-documents'
        ).json()
        assert sources[0]['document_id'] == draft['document_id']

        case_response = client.post(
            f'/api/documents?name=題目.pdf&ruleset_id={ruleset["id"]}',
            content=b'question-pdf',
            headers={'Content-Type': 'application/pdf'},
        )
        assert case_response.status_code == 200, case_response.text
        case = case_response.json()['case']
        assert case['ruleset_id'] == ruleset['id']
        assert case['locality'] == '測試市甲區'
        assert case['subject_address'] == '測試市甲區幸福路1號'
        assert case['comparable_address'] == '測試市甲區和平路2號'
        assert case['factors'][0]['subject'] == '85'

        export = client.get(
            f'/api/cases/{case["id"]}/export/review-xlsx?revision={case["revision"]}'
        )
        assert export.status_code == 200, export.text
        workbook = load_workbook(io.BytesIO(export.content))
        assert workbook.sheetnames == ['案件摘要', '區域因素', '個別因素', '審查結果']
        assert workbook['案件摘要']['B9'].value == '測試市甲區幸福路1號'
        assert workbook['區域因素']['A2'].value == '測試數值'
