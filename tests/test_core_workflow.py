"""Independent synthetic end-to-end oracles; no AWS or official case data."""
import io
import json
import zipfile
from pathlib import Path
import pytest
import httpx
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from app.interfaces.http import create_app
from app.infrastructure.settings import Settings
from app.infrastructure.external import OfficialEvidenceAdapter, measure
from app.application.ports import RevisionConflict
from app.domain.models import Case


def synthetic_rules():
    return dict(name='合成雙因素', version='test-1', locality='合成區', land_use='合成用地', source='tests/test_core_workflow.py 固定人工答案',
        direction='列：比準地；欄：比較標的；矩陣值為百分點',
        rules=[dict(id='width', name='寬度', group='道路', unit='m', scope='individual', source_page=1,
                    bands=[dict(label='窄', low=0, high=10), dict(label='寬', low=10, high=None)], matrix=[[0,-2],[2,0]]),
               dict(id='region', name='區域', group='區域', unit='', scope='regional', source_page=1,
                    bands=[dict(label='甲', values=['甲']), dict(label='乙', values=['乙'])], matrix=[[0,-1],[1,0]])])


def workbook():
    book = Workbook(); s = book.active; s.title='勘查'
    for key, value in {'A1':'表3 地價區段勘查表','A3':'年期','D11':'主要道路','J11':12,'F12':8,'H6':0,'H7':'無','H8':'不適用','J12':'=J11+1'}.items():
        s[key]=value
    stream=io.BytesIO();book.save(stream);return stream.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('SEED_EXAMPLES','false')
    monkeypatch.delenv('APP_ACCESS_TOKEN', raising=False)
    class Pdf:
        def read(self,data):
            if not data.startswith(b'%PDF-'):raise ValueError('不是 PDF')
            return [dict(page=1,text='合成勘查表 寬度 12 8',method='mock-ocr')]
    with TestClient(create_app(Settings(data_dir=tmp_path, ai_enabled=False),pdf=Pdf())) as client:
        yield client


def prepared(client, confirm=True):
    draft=client.post('/api/rulesets',json=synthetic_rules()).json()
    published=client.post('/api/rulesets/'+draft['id']+'/publish',json=dict(reason='合成驗收，人工核對固定矩陣',source_confirmed=True,matrix_confirmed=True,valid_from='2026-01-01',valid_to='2026-12-31')).json()
    body=dict(title='合成完整案件',case_number='SYNTHETIC-001',locality='合成區',land_use='合成用地',valuation_date='2026-09-01',
              subject_name='合成比準地',comparable_name='合成比較標的',subject_section='S1',comparable_section='S2',ruleset_id=published['id'],
              factors=[dict(id='width',subject='12',comparable='8',entered_rate=2),dict(id='region',subject='甲',comparable='甲',subject_grade='甲',comparable_grade='甲',entered_rate=0)],
              totals=dict(normal_price=100, time_rate=0, adjusted_price=100,regional_detail=0,regional_carried=0,individual=2,absolute=2,trial_price=102,weight=100))
    p=client.post('/api/cases',json=body).json()
    if confirm:
        for f in p['case']['factors']:f['confirmed']=True
        p['case']['totals_confirmed']=True
        p=client.put('/api/cases/'+p['case']['id'],json=p['case']).json()
    return p


def test_normal_upload_confirm_review_export(client):
    p=prepared(client)
    assert p['review']['complete'] is True
    assert p['review']['computed']['individual']==2 # Independent fixed golden value.
    cid=p['case']['id']
    response=client.post(f'/api/cases/{cid}/documents',params=dict(revision=p['case']['revision'],name='misleading-name.bin'),content=workbook())
    assert response.status_code==200,response.text
    uploaded=response.json()
    assert uploaded['parsed']['pages'][0]['kind']=='survey'
    assert uploaded['review']['complete'] is False
    document_id=uploaded['document_id']
    doc=client.get('/api/documents/'+document_id).json()
    cells=doc['pages'][0]['cells']
    assert cells['H6']['value']==0 and cells['H7']['presence']=='none' and cells['H8']['presence']=='not_applicable'
    assert cells['J12']['formula']=='=J11+1' and cells['J12']['value'] is None
    request=dict(revision=uploaded['case']['revision'],document_id=document_id,sheet='勘查',cell='J11',target='factors.width.subject')
    p=client.post(f'/api/cases/{cid}/apply-cell',json=request).json()
    assert p['case']['factors'][0]['evidence']['cell']=='J11'
    assert not p['case']['factors'][0]['confirmed']
    request.update(revision=p['case']['revision'],cell='J12')
    assert client.post(f'/api/cases/{cid}/apply-cell',json=request).status_code==400
    for f in p['case']['factors']:f['confirmed']=True
    p['case']['totals_confirmed']=True
    p=client.put(f'/api/cases/{cid}',json=p['case']).json()
    assert p['review']['complete']
    args=dict(revision=p['case']['revision'],run_id=p['run']['id'])
    pdf=client.get(f'/api/cases/{cid}/artifacts/pdf',params=args)
    assert pdf.status_code==200,pdf.text[:200] if pdf.status_code!=200 else ''
    text=''.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages)
    assert '合成完整案件' in text and 'J11' in text and '全部檢核通過' in text
    bundle=client.get(f'/api/cases/{cid}/artifacts/bundle',params=args)
    with zipfile.ZipFile(io.BytesIO(bundle.content)) as z:
        snapshot=json.loads(z.read('snapshot.json'))
        assert snapshot['run']['case_revision']==p['case']['revision']
        wb=load_workbook(io.BytesIO(z.read('confirmed-forms.xlsx')))
        assert wb['已確認勘查資料']['B4'].value=='12'
        assert wb['已確認勘查資料']['D4'].value==2


def test_error_missing_and_human_decisions_do_not_override_verifier(client):
    p=prepared(client);case=p['case'];cid=case['id']
    case['factors'][0]['entered_rate']=9
    p=client.put(f'/api/cases/{cid}',json=case).json()
    assert not p['case']['factors'][0]['confirmed']
    p['case']['factors'][0]['confirmed']=True;p['case']['totals_confirmed']=True
    p=client.put(f'/api/cases/{cid}',json=p['case']).json()
    check=next(c for c in p['review']['checks'] if c['id']=='width')
    assert check['status']=='error' and check['actual']==9 and check['expected']==2
    decision=dict(revision=p['case']['revision'],run_id=p['run']['id'],check_id='width',decision='reject',reason='合成拒絕示範',operation_id='synthetic-operation-1')
    first=client.post(f'/api/cases/{cid}/decisions',json=decision)
    assert first.status_code==200,first.text
    assert client.post(f'/api/cases/{cid}/decisions',json=decision).json()==first.json()
    p=client.get(f'/api/cases/{cid}').json()
    assert len(p['dispositions'])==1
    assert next(c for c in p['review']['checks'] if c['id']=='width')['status']=='error'
    p['case']['factors'][0]['subject']=None
    p=client.put(f'/api/cases/{cid}',json=p['case']).json()
    p['case']['factors'][0]['confirmed']=True
    p=client.put(f'/api/cases/{cid}',json=p['case']).json()
    assert p['review']['computed']['individual'] is None
    assert next(c for c in p['review']['checks'] if c['id']=='width')['status']=='missing'
    assert next(c for c in p['review']['checks'] if c['id']=='region')['status']=='pass'


def test_old_task_and_export_cannot_overwrite_latest_case(client):
    p=prepared(client);cid=p['case']['id'];workflow=client.app.state.service.workflow
    old=workflow.run(cid,p['case']['revision'])
    p['case']['factors'][0]['subject']='13'
    newer=client.put(f'/api/cases/{cid}',json=p['case']).json()
    with pytest.raises(RevisionConflict):workflow.repo.save_run(old)
    assert client.get(f'/api/cases/{cid}/artifacts/pdf',params=dict(revision=old['case_revision'],run_id=old['id'])).status_code==409
    assert client.get(f'/api/cases/{cid}/artifacts/pdf',params=dict(revision=newer['case']['revision'],run_id=old['id'])).status_code==409
    other=prepared(client)
    assert client.get(f"/api/cases/{other['case']['id']}/artifacts/pdf",params=dict(revision=other['case']['revision'],run_id=old['id'])).status_code==404
    original=workflow.external.query
    def racing(query):
        changed=workflow.repo.get_case(cid);changed.notes='concurrent'
        workflow.service.save_case(changed)
        return original(dict(mode='mock',scenario='timeout'))
    workflow.external.query=racing
    assert client.post(f'/api/cases/{cid}/external',json=dict(revision=newer['case']['revision'],factor_id='width',query={'mode':'mock'})).status_code==409
    assert workflow.repo.external_for(cid,newer['case']['revision'])==[]


def test_rule_publication_never_changes_bound_case(client):
    p=prepared(client);cid=p['case']['id'];old_id=p['case']['ruleset_id']
    draft=synthetic_rules();draft['approval_state']='published';draft['version']='test-2';draft['rules'][0]['matrix'][1][0]=3
    created=client.post('/api/rulesets',json=draft).json()
    assert created['approval_state']=='draft'
    published=client.post('/api/rulesets/'+created['id']+'/publish',json=dict(reason='核對新版',source_confirmed=True,matrix_confirmed=True,valid_from='2026-01-01',valid_to='2026-12-31')).json()
    unchanged=client.get(f'/api/cases/{cid}').json()
    assert unchanged['run']['id']==p['run']['id'] and unchanged['case']['ruleset_id']==old_id
    p['case']['ruleset_id']=published['id']
    changed=client.put(f'/api/cases/{cid}',json=p['case']).json()
    assert not any(f['confirmed'] for f in changed['case']['factors'])
    assert changed['run']['rules_hash']!=p['run']['rules_hash']
    assert client.app.state.service.repository.get_rules(created['id'])['approval_state']=='draft'


def test_external_failure_and_geometry_are_separate_from_table_checks(client):
    p=prepared(client);cid=p['case']['id']
    def denied(request):return httpx.Response(403)
    client.app.state.service.workflow.external=OfficialEvidenceAdapter(httpx.Client(transport=httpx.MockTransport(denied)))
    response=client.post(f'/api/cases/{cid}/external',json=dict(revision=p['case']['revision'],factor_id='width',query={'mode':'live','name':'公園'}))
    assert response.json()['status']=='unauthorized'
    newer=client.get(f'/api/cases/{cid}').json()
    assert newer['review']['complete'] is False
    assert next(c for c in newer['review']['checks'] if c['id']=='width')['status']=='pass'
    assert client.get(f'/api/cases/{cid}/artifacts/pdf',params=dict(revision=p['case']['revision'],run_id=p['run']['id'])).status_code==409
    result=measure(dict(method='straight_line',crs='EPSG:3826',start={'role':'survey_point','coordinates':[300000,2700000]},end={'role':'entrance','coordinates':[300003,2700004]}))
    assert result['distance_m']==5
    with pytest.raises(ValueError):measure(dict(method='walking',crs='EPSG:4326'))
    with pytest.raises(ValueError):measure(dict(method='straight_line',crs='EPSG:4326',start={'role':'center','coordinates':[121,25]}))


def test_unknown_files_parse_failure_and_network_access(client):
    p=prepared(client);cid=p['case']['id'];revision=p['case']['revision']
    result=client.post(f'/api/cases/{cid}/documents',params=dict(revision=revision,name='fake.pdf'),content=b'not a document')
    assert result.status_code==400
    assert client.get(f'/api/cases/{cid}').json()['case']['revision']==revision
    assert client.get(f'/api/cases/{cid}',headers={'Host':'evil.example'}).status_code==403
    with TestClient(client.app,client=('203.0.113.1',50000)) as remote:
        assert remote.get(f'/api/cases/{cid}').status_code==403


def test_three_comparisons_have_independent_checks_sources_weights_and_export(client):
    p=prepared(client);case=p['case'];cid=case['id']
    rules=synthetic_rules();rules.update(aggregation_formula='weighted-trial-v1',aggregation_source='合成三標的固定驗收公式：加權至元')
    draft=client.post('/api/rulesets',json=rules).json()
    published=client.post('/api/rulesets/'+draft['id']+'/publish',json=dict(reason='合成三標的公式核對',source_confirmed=True,matrix_confirmed=True,valid_from='2026-01-01',valid_to='2026-12-31')).json()
    case['ruleset_id']=published['id']
    case['totals']['weight']=40
    case['additional_comparisons']=[]
    for i in (2,3):
        factors=json.loads(json.dumps(case['factors']))
        factors[0]['comparable']='12';factors[0]['entered_rate']=0
        case['additional_comparisons'].append(dict(id=f'c{i}',name=f'比較{i}',section=f'S{i+1}',factors=factors,
            totals=dict(case['totals'],weight=30,individual=0,absolute=0,trial_price=100),totals_confirmed=True))
    p=client.put(f'/api/cases/{cid}',json=case).json()
    assert not p['review']['complete']
    for f in p['case']['factors']:f['confirmed']=True
    p['case']['totals_confirmed']=True
    for c in p['case']['additional_comparisons']:
        for f in c['factors']:f['confirmed']=True
        c['totals_confirmed']=True
    p=client.put(f'/api/cases/{cid}',json=p['case']).json()
    assert p['review']['complete'],p['review']
    assert p['review']['aggregate']=='101' # (102*.4 + 100*.3 + 100*.3) rounded independently.
    assert [c['computed']['individual'] for c in p['review']['comparisons']]==[2,0,0]
    exported=client.get(f'/api/cases/{cid}/artifacts/xlsx',params=dict(revision=p['case']['revision'],run_id=p['run']['id']))
    book=load_workbook(io.BytesIO(exported.content))
    assert [book['已確認勘查資料'].cell(row,7).value for row in (4,6,8)]==['primary','c2','c3']
    assert client.get(f'/api/cases/{cid}/export/forms').status_code==400
    p['case']['additional_comparisons'][0]['factors'][0]['entered_rate']=8
    changed=client.put(f'/api/cases/{cid}',json=p['case']).json()
    assert changed['case']['factors'][0]['confirmed'] is True
    assert changed['case']['additional_comparisons'][0]['factors'][0]['confirmed'] is False
    assert changed['case']['additional_comparisons'][1]['factors'][0]['confirmed'] is True
    assert changed['review']['aggregate'] is None


@pytest.mark.parametrize('status,expected',[(401,'unauthorized'),(403,'unauthorized'),(500,'unavailable'),(200,'no_match')])
def test_external_error_taxonomy(status,expected):
    adapter=OfficialEvidenceAdapter(httpx.Client(transport=httpx.MockTransport(lambda request:httpx.Response(status,json=[]))))
    result=adapter.query(dict(mode='live',name='合成不存在'))
    assert result['status']==expected and result['check_status']!='pass'


def test_late_background_calculation_and_new_evidence_are_rejected(client,monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    import app.application.workflow as module
    p=prepared(client);cid=p['case']['id'];workflow=client.app.state.service.workflow
    original=module.calculate;started=Event();release=Event()
    def slow(*args):
        started.set();assert release.wait(5);return original(*args)
    monkeypatch.setattr(module,'calculate',slow)
    with ThreadPoolExecutor(max_workers=1) as pool:
        pending=pool.submit(workflow.run,cid,p['case']['revision'])
        assert started.wait(5)
        changed=workflow.repo.get_case(cid);changed.notes='較新輸入'
        workflow.repo.save_case(changed,'合成競爭修改')
        release.set()
        with pytest.raises(RevisionConflict):pending.result(timeout=5)
    monkeypatch.setattr(module,'calculate',original)
    case=workflow.repo.get_case(cid);old=workflow.run(cid,case.revision)
    workflow.lookup(cid,case.revision,'width',dict(mode='mock',scenario='timeout'))
    with pytest.raises(RevisionConflict):workflow.repo.save_run(old)


def test_draft_period_and_unknown_direction_cannot_pass(client):
    p=prepared(client);case=p['case']
    case['valuation_date']='2027-01-01'
    p=client.put('/api/cases/'+case['id'],json=case).json()
    assert p['review']['complete'] is False
    assert any('適用期間' in r['message'] for r in p['review']['checks'])
    raw=synthetic_rules();raw['direction']='列：比較標的；欄：比準地'
    assert client.post('/api/rulesets',json=raw).status_code==400


def test_pdf_long_human_explanation_is_renderable(client):
    p=prepared(client);case=p['case'];case['factors'][0]['note']='合成長篇核對理由。'*290
    case['factors'][0]['entered_rate']=9
    p=client.put('/api/cases/'+case['id'],json=case).json()
    p['case']['factors'][0]['confirmed']=True
    p=client.put('/api/cases/'+case['id'],json=p['case']).json()
    result=client.get(f"/api/cases/{case['id']}/artifacts/pdf",params=dict(revision=p['case']['revision'],run_id=p['run']['id']))
    assert result.status_code==200
    assert len(PdfReader(io.BytesIO(result.content)).pages)>0


def test_agent_can_plan_and_query_official_adapter_without_changing_case(client):
    from app.application.agent_contracts import AgentTurn,ToolCall
    from app.application.rag_contracts import AnswerDraft
    p=prepared(client);cid=p['case']['id'];service=client.app.state.service
    service.workflow.external=OfficialEvidenceAdapter(httpx.Client(transport=httpx.MockTransport(lambda request:httpx.Response(403))))
    class Model:
        def next_turn(self,context,history,tools):
            assert {'plan_checks','lookup_facility'} <= {t['name'] for t in tools}
            if not history:
                return AgentTurn(calls=(ToolCall(id='p',name='plan_checks',arguments={}),ToolCall(id='e',name='lookup_facility',arguments={'rule_id':'width','name':'公園'})))
            if len(history)==1:
                assert history[0]['results'][1]['data']['status']=='unauthorized'
                return AgentTurn(calls=(ToolCall(id='r',name='review_case',arguments={}),))
            return AgentTurn(answer=AnswerDraft(statements=(),insufficient_evidence=True))
    service.rag.agent.model=Model()
    result=client.post(f'/api/cases/{cid}/agent-evidence',json=dict(revision=p['case']['revision'],question='規劃審查並查證公園',cloud_data_approved=True))
    assert result.status_code==200,result.text
    assert result.json()['review']['complete'] is False
    assert result.json()['external_observations'][0]['persisted'] is False
    assert service.repository.external_for(cid,p['case']['revision'])==[]
    assert client.get(f'/api/cases/{cid}').json()['case']==p['case']
