import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient

from app.application.export_contracts import ExportUnavailable
from app.domain.sample import sample_case
from app.domain.rules import default_rules
from app.domain.engine import review
from app.infrastructure.form_exports import TemplateFormRenderer, TEMPLATES
from app.infrastructure.settings import Settings
from app.interfaces.http import create_app


@pytest.fixture
def templates(tmp_path):
    for number, (_, name, _, _) in TEMPLATES.items():
        wb = Workbook()
        wb.active.title = name
        wb.active['A1'] = f'表{number} 合成模板'
        wb.active.merge_cells('A1:C1')
        if number == '4':
            wb.active.merge_cells('G29:J29')
        wb.active['A2'] = '=HYPERLINK("https://example.invalid","sample")'
        hidden = wb.create_sheet('隱藏範例')
        hidden.sheet_state = 'hidden'
        hidden['A1'] = '不得輸出的舊價格'
        wb.save(tmp_path / f'表{number}合成.xlsx')
    return tmp_path


def render(renderer, kind, case=None):
    case = case or sample_case(False)
    rules = default_rules()
    return renderer.render(case, review(case, rules), rules, kind, '2026-09-12T00:00:00Z')


def test_workbooks_fill_separate_forms_keep_zero_and_literal_text(templates):
    case = sample_case(False)
    case.subject_name = '=HYPERLINK("https://example.invalid")'
    case.totals.time_rate = 0
    source_bytes = {p: p.read_bytes() for p in templates.glob('*.xlsx')}
    for number in ('3', '4', '5'):
        artifact = render(TemplateFormRenderer(templates), f'table{number}-xlsx', case)
        wb = load_workbook(io.BytesIO(artifact.data))
        assert '隱藏範例' in wb.sheetnames
        assert '填值與審核明細' in wb.sheetnames
        if number == '3':
            assert wb['表3區段勘查表']['A2'].data_type == 'f'
        assert 'A1:C1' in str(wb.active.merged_cells)
        if number == '3':
            assert len(wb.worksheets) == 3
            assert case.subject_section == wb['表3區段勘查表']['G3'].value
        if number == '4':
            assert wb.active['D4'].value == case.subject_name
            assert wb.active['D4'].data_type == 's'
            assert wb.active['J6'].value == 0
            assert wb.active['K4'].value == '未提供標的'
            assert wb.active['G31'].value == case.totals.trial_price
        if number == '5':
            assert wb.active['E42'].value == '基準不適用'
    assert all(p.read_bytes() == data for p, data in source_bytes.items())


def test_legacy_openpyxl_defined_name_collection_is_supported(templates, monkeypatch):
    """The deployed Anaconda image uses openpyxl 3.0's DefinedNameList."""
    import openpyxl.workbook.defined_name as defined_name_module
    legacy = defined_name_module.DefinedNameList
    assert hasattr(legacy(), 'definedName')
    artifact = render(TemplateFormRenderer(templates), 'table3-xlsx')
    assert artifact.data.startswith(b'PK')


@pytest.mark.parametrize('kind', ['report-pdf', 'table3-pdf', 'table4-pdf', 'table5-pdf'])
def test_pdf_exports_are_removed(templates, tmp_path, kind):
    with pytest.raises(KeyError):
        render(TemplateFormRenderer(templates), kind)
    app = create_app(Settings(data_dir=tmp_path / 'data', form_template_dir=templates, ai_enabled=False))
    with TestClient(app) as client:
        case = client.get('/api/cases').json()[0]
        response = client.get(f'/api/cases/{case["id"]}/export/{kind}?revision=1')
        assert response.status_code == 404


def test_exports_require_revision_and_do_not_mutate_case(templates, tmp_path):
    app = create_app(Settings(data_dir=tmp_path / 'data', form_template_dir=templates, ai_enabled=False))
    with TestClient(app) as client:
        case = client.get('/api/cases').json()[0]
        cid = case['id']
        current = client.get(f'/api/cases/{cid}').json()['case']
        url = f'/api/cases/{cid}/export/table4-xlsx'
        assert client.get(url).status_code == 422
        assert client.get(url + '?revision=-1').status_code == 409
        response = client.get(url + f'?revision={current["revision"]}')
        assert response.status_code == 200
        assert response.headers['x-case-revision'] == str(current['revision'])
        assert load_workbook(io.BytesIO(response.content)).active['D4'].value == current['subject_name']
        assert client.get(f'/api/cases/{cid}').json()['case'] == current
        app.state.service.renderer = TemplateFormRenderer(tmp_path / 'missing')
        assert client.get(url + f'?revision={current["revision"]}').status_code == 503


def test_revision_change_during_render_is_rejected(templates, tmp_path):
    app = create_app(Settings(data_dir=tmp_path / 'data', form_template_dir=templates, ai_enabled=False))
    with TestClient(app) as client:
        service = app.state.service
        case = service.repository.list_cases()[0][0]
        original = service.renderer
        class ConcurrentRenderer:
            def render(self, snapshot, *args):
                artifact = original.render(snapshot, *args)
                service.repository.save_case(snapshot, 'concurrent update')
                return artifact
        service.renderer = ConcurrentRenderer()
        response = client.get(f'/api/cases/{case.id}/export/table4-xlsx?revision={case.revision}')
        assert response.status_code == 409
